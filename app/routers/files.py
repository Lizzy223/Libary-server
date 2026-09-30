"""CSV / Excel reading shared by the member and catalogue importers."""
import csv
import io

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

MAX_BYTES = 10 * 1024 * 1024


async def read_rows(file: UploadFile) -> list[dict]:
    raw = await file.read()
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "File is too large (10 MB max).")
    name = (file.filename or "").lower()
    try:
        if name.endswith((".xlsx", ".xlsm")):
            ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            header = [str(h or "").strip().lower() for h in rows[0]]
            return [dict(zip(header, [("" if v is None else str(v).strip()) for v in r])) for r in rows[1:]
                    if any(v not in (None, "") for v in r)]
        text = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in reader
                if any((v or "").strip() for v in row.values())]
    except HTTPException:
        raise
    except Exception:  # noqa: BLE001
        raise HTTPException(400, "Could not read the file. Use the CSV template or a plain .xlsx sheet.") from None


def csv_response_body(columns: list[str], rows: list[list]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    writer.writerows(rows)
    return buf.getvalue()
